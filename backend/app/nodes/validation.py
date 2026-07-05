"""
Validation node: deterministic, no LLM call. The OfficeSkill node already has the
model self-verify (render to PDF, visually inspect a page) as part of following the
skill's own documented workflow -- this node is a second, independent, non-LLM check
that the file is actually well-formed and contains the planned content, reusing the
skill repo's own OOXML scripts rather than reimplementing format parsing.
"""
import re
import subprocess
import tempfile
from pathlib import Path

from app.state import DocGenState

REPO_ROOT = Path(__file__).resolve().parents[3]


def validation_node(state: DocGenState) -> dict:
    if state.get("error"):
        return {"validation": {"ok": False, "details": state["error"]}}

    fmt = state["format"]
    path = Path(state["file_path"])
    checker = {
        "docx": _check_docx,
        "pptx": _check_pptx,
        "xlsx": _check_xlsx,
        "pdf": _check_pdf,
    }[fmt]
    return {"validation": checker(path, fmt, state["content_plan"])}


def _unpack(path: Path, fmt: str, tmp: str):
    """Returns the unpack.py subprocess result; caller checks returncode."""
    unpack_script = REPO_ROOT / "mnt" / "skills" / "public" / fmt / "ooxml" / "scripts" / "unpack.py"
    return subprocess.run(
        ["python3", str(unpack_script), str(path), tmp],
        capture_output=True,
        text=True,
    )


def _check_docx(path: Path, fmt: str, content_plan: dict) -> dict:
    with tempfile.TemporaryDirectory() as tmp:
        result = _unpack(path, fmt, tmp)
        if result.returncode != 0:
            return {"ok": False, "details": f"unpack.py failed: {result.stderr}"}

        # A title landing in a running header/footer (word/header*.xml,
        # word/footer*.xml) rather than the body is a legitimate professional layout
        # choice, not a content defect -- observed for real in testing. Check all of
        # them, not just document.xml.
        xml_globs = ["word/document.xml", "word/header*.xml", "word/footer*.xml"]
        matches = [p for g in xml_globs for p in Path(tmp).glob(g)]
        text = "\n".join(p.read_text(errors="ignore") for p in matches)
        title = content_plan.get("title", "")
        missing = _missing_words(title, text)
        if missing:
            return {"ok": False, "details": f"Title words not found in generated docx content: {sorted(missing)}"}
        return {"ok": True, "details": "Valid docx, title confirmed present"}


def _check_pptx(path: Path, fmt: str, content_plan: dict) -> dict:
    """
    Unlike docx, pptx titles are NOT checked for word overlap with the content plan.
    html2pptx.md (the skill's own design guidance) explicitly instructs the model to
    invent creative branding/product names and taglines rather than restate a generic
    working title verbatim -- e.g. a content-plan title of "SaaS Product Pitch Deck"
    legitimately became "FlowSync -- Where Teams Move as One" on the actual slide, with
    zero word overlap. That's the skill working as documented, not a defect. Checking
    for literal title fidelity here would mean fighting the skill's own documented
    creative process, which is exactly what this whole architecture is built to avoid.
    Instead: confirm the file unpacks, has at least one slide, and each slide actually
    has non-empty text content (catches a genuinely broken/empty deck without policing
    wording or slide count the skill is deliberately free to reinvent).

    An earlier version of this check also required len(slides) >= count(headings) + 1,
    assuming a roughly 1:1 mapping from content-plan sections to slides. That's also a
    wrong assumption: html2pptx.md tells the model to keep slides brief/uncluttered, so
    it may legitimately consolidate multiple sections onto one slide. Observed live: a
    4-heading content plan for an explicitly-requested "4-slide deck" correctly produced
    a well-formed, complete 4-slide deck (title, problem, solution, CTA) that this
    check flagged as "expected >= 5 slides" -- a false failure on a genuinely correct
    result. Slide *count* isn't ours to dictate any more than title wording is.
    """
    with tempfile.TemporaryDirectory() as tmp:
        result = _unpack(path, fmt, tmp)
        if result.returncode != 0:
            return {"ok": False, "details": f"unpack.py failed: {result.stderr}"}

        slide_files = sorted(Path(tmp).glob("ppt/slides/slide*.xml"))
        if not slide_files:
            return {"ok": False, "details": "No slides found in generated pptx"}

        empty_slides = [
            s.name for s in slide_files
            if not re.search(r"<a:t>[^<]*[^\s<][^<]*</a:t>", s.read_text(errors="ignore"))
        ]
        if empty_slides:
            return {"ok": False, "details": f"Slides with no text content: {empty_slides}"}

        return {"ok": True, "details": f"Valid pptx, {len(slide_files)} non-empty slides"}


def _missing_words(title: str, text: str) -> set:
    """
    Word-level, case-insensitive containment instead of exact-substring matching.
    Observed live: a title can legitimately be rendered in all-caps, placed in a
    header/footer rather than the body, or split across separate text runs/text boxes
    (e.g. a pptx title + subtitle rendered as two elements, joined by an em/en dash in
    the content plan but never appearing as one contiguous string in the XML). None of
    these are content defects -- checking that the title's words all appear somewhere,
    rather than requiring one exact contiguous phrase, tolerates all of them while
    still catching a genuinely missing/wrong title. Words of length <= 2 are skipped
    to avoid noise from articles/prepositions.
    """
    title_words = {w for w in re.findall(r"[a-z0-9]+", title.lower()) if len(w) > 2}
    text_words = set(re.findall(r"[a-z0-9]+", text.lower()))
    return title_words - text_words


def _check_xlsx(path: Path, fmt: str, content_plan: dict) -> dict:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        return {"ok": False, "details": f"openpyxl could not open file: {e}"}
    if not wb.sheetnames:
        return {"ok": False, "details": "Workbook has no sheets"}
    return {"ok": True, "details": f"Valid xlsx with sheets: {wb.sheetnames}"}


def _check_pdf(path: Path, fmt: str, content_plan: dict) -> dict:
    header = path.read_bytes()[:5]
    if header != b"%PDF-":
        return {"ok": False, "details": f"Missing %PDF- header, got {header!r}"}
    return {"ok": True, "details": "Valid PDF header"}
